#!/usr/bin/env python3
"""Build report figures and an editable DOCX from archived CSVs and Markdown.
Dependencies: numpy, matplotlib, python-docx. No robot access or data mutation.
"""
import csv
import json
import re
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/yaw_identification/submission'
FIG=OUT/'figures';FIG.mkdir(parents=True,exist_ok=True)
plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,
                     'axes.grid':True,'grid.alpha':.2,'savefig.dpi':220})


def raw(name):
    with (ROOT/'data'/name).open() as f:
        reader=csv.DictReader(f);rows=[]
        for r in reader:
            try:
                converted={k:float(v) for k,v in r.items()}
            except (TypeError,ValueError):continue
            rows.append(converted)
    return {k:np.array([r[k] for r in rows]) for k in rows[0]}


def predictions(folder):
    with (ROOT/'docs/yaw_identification'/folder/'prediction.csv').open() as f:
        rows=list(csv.DictReader(f))
    return {k:np.array([float(r[k]) for r in rows]) for k in rows[0]}


def finish(fig,name):
    fig.tight_layout()
    fig.savefig(FIG/(name+'.png'),bbox_inches='tight')
    fig.savefig(FIG/(name+'.svg'),bbox_inches='tight')
    plt.close(fig)


def plots():
    d=raw('yaw扫频测试_原代码_2026-09-22_09-00-27.csv')
    mask=d['/gimbal/yaw/excitation/state']==2
    t=d['/gimbal/sample_time_s'][mask];t=t-t[0]
    fig,ax=plt.subplots(3,1,figsize=(9,7),sharex=True)
    # Display every fifth original sample; all analysis metrics use full resolution.
    sl=slice(None,None,5)
    for key,label in [('excitation/velocity_rad_s','Reference'),('velocity_imu','Measured IMU')]:
        ax[0].plot(t[sl],d['/gimbal/yaw/'+key][mask][sl],label=label,lw=1)
    ax[0].set_ylabel('Yaw velocity (rad/s)');ax[0].legend(ncol=2)
    for key,label in [('control_torque','Command'),('torque','Current-derived feedback')]:
        ax[1].plot(t[sl],d['/gimbal/yaw/'+key][mask][sl],label=label,lw=.9,alpha=.85)
    ax[1].set_ylabel('Torque (N m)');ax[1].legend(ncol=2)
    ax[2].plot(t[sl],np.degrees(d['/gimbal/yaw/control_angle_error'][mask][sl]),color='#b45309',lw=1)
    ax[2].set_ylabel('Angle error (deg)');ax[2].set_xlabel('Time since sweep start (s)')
    fig.suptitle('Original controller: 0.1 to 5 Hz sweep (09:00:27)',y=1.01)
    finish(fig,'01_sweep')

    sweep=predictions('validation_2026-09-22_09-01-51')
    manual=predictions('validation_manual_with_5hz_model')
    fig,ax=plt.subplots(3,1,figsize=(9,8))
    t=sweep['time_s']-sweep['time_s'][0]
    for key,label in [('measured_velocity_rad_s','Measured'),('sweep_full','Frozen model')]:
        ax[0].plot(t[::5],sweep[key][::5],label=label,lw=1)
        m=(t>=10)&(t<=13)
        ax[1].plot(t[m],sweep[key][m],label=label,lw=1.2)
    ax[0].set_title('Independent repeated sweep: RMSE 0.181 rad/s; R2 0.955')
    ax[1].set_title('Same sweep: 10-13 s detail')
    for seg in np.unique(manual['segment']):
        m=manual['segment']==seg
        for key,label in [('measured_velocity_rad_s','Measured'),('sweep_full','Frozen model')]:
            ax[2].plot(manual['time_s'][m][::5],manual[key][m][::5],
                       color='#1f77b4' if key=='measured_velocity_rad_s' else '#ff7f0e',
                       label=label if seg==1 else None,lw=1)
    ax[2].set_title('Manual operation: RMSE 1.523 rad/s; R2 0.583')
    for axis in ax:
        axis.set_ylabel('Yaw velocity (rad/s)');axis.legend(loc='upper right',ncol=2)
    ax[0].set_xlabel('Time since sweep start (s)');ax[1].set_xlabel('Time since sweep start (s)')
    ax[2].set_xlabel('Recorded sample time (s); separate initialization per enabled segment')
    finish(fig,'02_validation')

    rows=json.loads((ROOT/'docs/yaw_identification/frequency_response_2026-09-22/frequency_response.json').read_text())
    fig,axes=plt.subplots(2,2,figsize=(9,6.4))
    specs=[('velocity_gain','Velocity amplitude ratio'),('velocity_lag_deg','Velocity lag (deg)'),
           ('torque_gain','Feedback / command torque amplitude'),('torque_lag_deg','Torque lag (deg)')]
    for index,source in enumerate(sorted({r['source'] for r in rows})):
        r=[x for x in rows if x['source']==source]
        f=[(x['frequency_low_hz']+x['frequency_high_hz'])/2 for x in r]
        for axis,(key,label) in zip(axes.flat,specs):
            axis.plot(f,[x[key] for x in r],marker='o' if index==0 else 'x',
                      ls='-' if index==0 else '--',label=f'Run {index+1}',lw=1.3)
            axis.set_ylabel(label);axis.set_xlabel('Band center (Hz)')
    axes[0,0].axhline(1/np.sqrt(2),color='gray',ls=':',lw=1,label='-3 dB')
    for axis in axes.flat:axis.legend()
    fig.suptitle('Local chirp estimates; fade-in/out excluded',y=1.01)
    finish(fig,'03_frequency_response')


def inline(paragraph,text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]*`)',text):
        if part.startswith('**'):
            run=paragraph.add_run(part[2:-2]);run.bold=True
        elif part.startswith('`'):
            run=paragraph.add_run(part[1:-1]);run.font.name='Consolas';run.font.size=Pt(9)
        else:paragraph.add_run(part)


def docx():
    path=OUT/'云台系统辨识阶段总结.md'
    lines=path.read_text().splitlines();doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(1.8);section.left_margin=section.right_margin=Cm(1.8)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','List Bullet']:
        style=doc.styles[name];style.font.name='Calibri'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.18
    footer=section.footer.paragraphs[0];footer.alignment=2
    footer.add_run('云台系统辨识 · 阶段总结  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    i=0;eq=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line=='$$':
            content=[]
            while i<len(lines) and lines[i].strip()!='$$':content.append(lines[i]);i+=1
            i+=1;eq+=1;formula=' '.join(content)
            fig=plt.figure(figsize=(7,.65));fig.text(.5,.5,'$'+formula+'$',ha='center',va='center',fontsize=17)
            image=FIG/f'equation_{eq:02d}.png';fig.savefig(image,bbox_inches='tight',pad_inches=.06,dpi=240);plt.close(fig)
            para=doc.add_paragraph();para.alignment=1;para.add_run().add_picture(str(image),width=Cm(14))
        elif line.startswith('```'):
            content=[]
            while i<len(lines) and not lines[i].startswith('```'):content.append(lines[i]);i+=1
            i+=1
            para=doc.add_paragraph();run=para.add_run('\n'.join(content));run.font.name='Consolas';run.font.size=Pt(8)
        elif line.startswith('|'):
            tablelines=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):tablelines.append(lines[i].strip());i+=1
            data=[[v.strip() for v in row.strip('|').split('|')] for row in tablelines]
            data=[row for row in data if not all(re.fullmatch(r'[:\- ]+',v) for v in row)]
            table=doc.add_table(rows=0,cols=len(data[0]));table.style='Light Shading Accent 1'
            for ri,row in enumerate(data):
                cells=table.add_row().cells
                for cell,value in zip(cells,row):
                    inline(cell.paragraphs[0],value)
                    for run in cell.paragraphs[0].runs:run.font.size=Pt(9)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                prop=table.rows[-1]._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
            doc.add_paragraph()
        elif line.startswith('!['):
            match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            doc.add_picture(str(OUT/match[2]),width=Cm(17.1))
            doc.paragraphs[-1].alignment=1
        elif line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));title=line[level:].strip()
            doc.add_heading(title,level=0 if level==1 else min(level-1,3))
        elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
        elif line.startswith('> '):inline(doc.add_paragraph(),line[2:])
        else:inline(doc.add_paragraph(),line)
    doc.core_properties.title='控制作业：步兵云台 yaw 轴系统辨识'
    doc.core_properties.author='谢元宏'
    target=OUT/'云台系统辨识阶段总结.docx';doc.save(target)
    check=Document(target)
    assert len(check.inline_shapes)==eq+3
    assert len(check.tables)>=8
    print(f'Created {target}; {len(check.tables)} tables, {len(check.inline_shapes)} figures/equations')

if __name__=='__main__':
    plots();docx()
